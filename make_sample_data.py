#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""生成合成“发货跟踪表”工作簿（无任何真实业务信息）。

模拟真实模板的复杂之处：
- 一个工作簿内多个供应商 sheet（列名写法不同，验证别名定位）
- 日期明细列位于“合同金额”与“剩余未排”之间（真实模板布局）
- 一个 hidden sheet（验证过滤）+ 一个损坏文件（验证失败隔离）
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).parent / "data" / "orders"

SHIP_DATES = [date(2024, 3, 10), date(2024, 3, 15), date(2024, 3, 20), date(2024, 3, 25)]


def sheet_rows(variant: int) -> list:
    """variant=0 用标准列名，variant=1 用别名变体，验证多别名定位。"""
    if variant == 0:
        spec_h, surplus_h, order_h, price_h = "规格型号", "剩余未排/kg", "采购合同号", "合同单价"
    else:
        spec_h, surplus_h, order_h, price_h = "品相", "剩余未排", "晶澳合同号", "承兑单价"
    rows = [[None] * 9 for _ in range(3)]
    rows[0] = [None] * 4 + SHIP_DATES[:3] + [None] * 2          # 日期行
    rows[1] = ["供应商", spec_h, "合同金额/元", "不含税电汇金额/元",
               None, None, None, surplus_h, order_h]             # 表头行
    sup = "NovaSilicon" if variant == 0 else "HelioMaterials"
    rows[2] = [sup, "dense small", 1000, None,
               30, 0, 20, 50, "PO-2024-031"]                     # 数据行（明细 30+20kg）
    rows.append([sup, "loose chunk", 800, None, 0, 15, 0, 60, "PO-2024-032"])
    rows.append(["合计"] + [None] * 8)                            # 表尾（定位 down border）
    return rows


def build(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "发货-NovaSilicon"
    for row in sheet_rows(0):
        ws.append(row)
    ws["J2"] = "收货基地"; ws["K2"] = price_h = "合同单价"
    ws["J3"] = "包头基地"; ws["J4"] = "曲靖基地"
    ws["K3"] = 68.5; ws["K4"] = 70.0

    ws2 = wb.create_sheet("发货-Helio")
    for row in sheet_rows(1):
        ws2.append(row)
    ws2["J2"] = "基地"; ws2["K2"] = "承兑单价"
    ws2["J3"] = "包头基地"; ws2["J4"] = "曲靖基地"
    ws2["K3"] = 66.0; ws2["K4"] = 69.5

    hidden = wb.create_sheet("发货-Old")
    hidden.append(["垃圾数据", None, None])
    hidden.sheet_state = "hidden"

    wb.save(path)


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    build(ROOT / "delivery_tracker_2024W11.xlsx")
    (ROOT / "corrupted.xlsx").write_bytes(b"not a real xlsx")
    print(f"sample data generated under {ROOT}")


if __name__ == "__main__":
    main()
